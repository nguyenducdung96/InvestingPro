"""Archive authority activity workbooks and extract only explicitly attributed NAP rows."""
import argparse
from concurrent.futures import ThreadPoolExecutor, as_completed
import datetime as dt
import hashlib
import io
import json
from pathlib import Path
import re
import time
import unicodedata
import urllib.parse
import urllib.request
import zipfile
from bs4 import BeautifulSoup
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'.local/nap'
BASE='https://cangvuhanghainghean.gov.vn/'
INDEX=BASE+'index.aspx?page=news&cat=2017'
VERSION='nap-authority-xlsx/1.0.0'

def clean(value):
    return re.sub(r'\s+',' ',unicodedata.normalize('NFC',str(value or ''))).strip()

def allowed(url):
    p=urllib.parse.urlsplit(url)
    return p.scheme=='https' and p.netloc=='cangvuhanghainghean.gov.vn' and not p.fragment

def fetch(url,body=None,limit=2_000_000):
    if not allowed(url):raise ValueError('Unexpected source URL')
    req=urllib.request.Request(url,data=body,headers={'User-Agent':'Mozilla/5.0','Content-Type':'application/x-www-form-urlencoded'})
    for attempt in range(3):
        try:
            with urllib.request.urlopen(req,timeout=40) as r:
                if not allowed(r.geturl()):raise ValueError('Unexpected redirect')
                raw=r.read(limit+1)
            if len(raw)>limit:raise ValueError('Source size limit')
            return raw
        except OSError:
            if attempt==2:raise
            time.sleep(2)

def archive(raw,suffix):
    digest=hashlib.sha256(raw).hexdigest()
    folder=OUT/'archive';folder.mkdir(parents=True,exist_ok=True)
    (folder/(digest+suffix)).write_bytes(raw)
    return digest

def publication(text):
    m=re.search(r'(\d{1,2})/(\d{1,2})/(\d{4}),\s*(\d{1,2}):(\d{2})',text)
    if not m:raise ValueError('Publication date missing')
    day=str(dt.date(int(m[3]),int(m[2]),int(m[1])))
    return day,day+'T'+m[4].zfill(2)+':'+m[5]+':00+07:00'

def discover(start,resume=False):
    found={};body=None;pages=[]
    for page in range(1,150):
        pointer=OUT/f'index-{page}.html'
        raw=pointer.read_bytes() if resume and pointer.exists() else fetch(INDEX,body)
        pointer.write_bytes(raw);digest=archive(raw,'.html')
        s=BeautifulSoup(raw,'html.parser')
        current=s.select_one('#ctl22_LabelCurrentPage')
        if not current or current.get_text(strip=True)!=str(page):raise ValueError('Index did not advance')
        dates=[]
        for a in s.select('a.titleNewsNormal[href]'):
            item=a.find_parent('li');date=item.select_one('.dateNews')
            if not date:raise ValueError('Dated report entry missing')
            day,published=publication(date.get_text(' ',strip=True));dates.append(day)
            url=urllib.parse.urljoin(BASE,a['href'])
            if not re.fullmatch(re.escape(BASE)+r'index\.aspx\?page=detail&id=\d+',url):raise ValueError('Detail link changed')
            if day>=start:found[url]={'id':int(urllib.parse.parse_qs(urllib.parse.urlsplit(url).query)['id'][0]),'url':url,'day':day,'published_at':published,'index_sha256':digest}
        if not dates:raise ValueError('No dated reports on index')
        pages.append({'page':page,'sha256':digest})
        print(json.dumps({'page':page,'through':min(dates),'reports':len(found)}),flush=True)
        if min(dates)<start:break
        fields={i['name']:i.get('value','') for i in s.select('input[type=hidden][name]')}
        fields.update({'ctl22$ImageButton_NavNext.x':'1','ctl22$ImageButton_NavNext.y':'1'})
        body=urllib.parse.urlencode(fields).encode()
    else:raise ValueError('History bound exhausted')
    result={'checked_at':dt.datetime.now(dt.timezone.utc).isoformat(),'source':INDEX,'pages':pages,'reports':list(found.values())}
    (OUT/'candidates.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    return result

def scope(label):
    text=clean(label).upper()
    if text.startswith('BẾN CẢNG') and 'CỬA LÒ' in text and 'NGHỆ TĨNH' in text:return 'nap_cua_lo'
    if text.startswith(('CẢNG BẾN THỦY','CẢNG BẾN THUỶ','BẾN CẢNG BẾN THỦY','BẾN CẢNG BẾN THUỶ')) and 'NGHỆ TĨNH' in text:return 'nap_ben_thuy'
    return None

def weight(cell):
    value=cell.value
    if isinstance(value,bool) or value is None:return None
    if isinstance(value,(float,int)):number=float(value)
    else:
        text=clean(value).replace(' ','')
        if not re.fullmatch(r'\d+(?:[.,]\d+)*',text):return None
        if ',' in text:text=text.replace('.','').replace(',','.')
        elif re.fullmatch(r'\d{1,3}(?:\.\d{3})+',text):text=text.replace('.','')
        try:number=float(text)
        except ValueError:return None
    return number if 0<number<=1_000_000 else None

def date_value(cell,report=False):
    value=cell.value
    if value is None or clean(value) in ('','/','-'):return None
    precise=False
    if isinstance(value,dt.datetime):
        parsed=value
        precise=bool(re.search(r'[hs]',re.sub(r'"[^"]*"|\\.','',cell.number_format),re.I))
    elif isinstance(value,dt.date):parsed=dt.datetime.combine(value,dt.time())
    elif isinstance(value,str):
        text=clean(value)
        if report:text=re.sub(r'^NGÀY\s+','',text,flags=re.I)
        m=re.fullmatch(r'(\d{1,2})[/.](\d{1,2})[/.](\d{4})(?:\s+(\d{1,2}):(\d{2}))?',text)
        if not m:raise ValueError('Invalid source date')
        parsed=dt.datetime(int(m[3]),int(m[2]),int(m[1]),int(m[4] or 0),int(m[5] or 0));precise=m[4] is not None
    else:raise ValueError('Unformatted event date')
    day=parsed.date().isoformat()
    return {'day':day,'time':parsed.strftime('%H:%M') if precise else None,'precision':'minute' if precise else 'day'}

def raw_cell(cell):
    v=cell.value
    return {'value':v.isoformat() if isinstance(v,(dt.date,dt.datetime)) else v,'format':cell.number_format,'type':cell.data_type}

def parse_workbook(raw,published_day):
    if len(raw)>500_000:raise ValueError('Workbook size limit')
    z=zipfile.ZipFile(io.BytesIO(raw))
    if sum(i.file_size for i in z.infolist())>10_000_000:raise ValueError('Expanded workbook size limit')
    w=load_workbook(io.BytesIO(raw),data_only=True)
    formulas=load_workbook(io.BytesIO(raw),data_only=False)
    if w.epoch.year!=1899:raise ValueError('Unsupported date system')
    candidates=[name for name in w.sheetnames if clean(name).casefold()=='daily report']
    if len(candidates)!=1:raise ValueError('Daily Report sheet missing or ambiguous')
    sheet=w[candidates[0]];f=formulas[candidates[0]]
    if 'BÁO CÁO HOẠT ĐỘNG TÀU THUYỀN' not in clean(sheet['A1'].value):raise ValueError('Workbook title changed')
    report=date_value(sheet['A2'],report=True)
    if not report:raise ValueError('Workbook report date missing')
    lag=(dt.date.fromisoformat(published_day)-dt.date.fromisoformat(report['day'])).days
    if lag<0 or lag>90:raise ValueError('Workbook report date outside publication window')
    if f['A2'].data_type=='f' and f['A2'].value!='=TODAY()':raise ValueError('Unexpected date formula')
    area=None;header=False;departure_header=False;headings=[];rows=[];rejected=[];schema_errors=[];valid_headers=0
    for cells in sheet.iter_rows(max_row=min(sheet.max_row,1500),max_col=25):
        n=cells[0].row;name=clean(cells[1].value)
        if re.match(r'^(?:BẾN |CẢNG |KHU NEO|THI CÔNG)',name):
            area=scope(name);header=False;headings.append({'row':n,'label':name,'scope':area});continue
        if 'Tên tàu' in name:
            header=bool(re.search(r'D\.?W\.?T',clean(cells[6].value),re.I) and 'đến' in clean(cells[7].value).lower() and 'đi' in clean(cells[8].value).lower())
            departure_header=bool(re.search(r'đi|departure',clean(cells[8].value),re.I))
            header=bool(re.search(r'D\.?W\.?T',clean(cells[6].value),re.I) and 'đến' in clean(cells[7].value).lower())
            if area and (not header or not departure_header):schema_errors.append({'scope':area,'row':n,'arrival_header_verified':header,'departure_header_verified':departure_header})
            if area and header:valid_headers+=1
            continue
        if not area or not header or not name or not re.fullmatch(r'\d+',clean(cells[0].value)):continue
        evidence={c.column_letter:raw_cell(c) for c in cells if c.value is not None}
        reason=None
        if any(f.cell(n,col).data_type=='f' for col in [2,7,8,9]):reason='formula_in_vessel_cells'
        if '+' in name:reason='multiple_vessels'
        try:arrival=date_value(cells[7])
        except (ValueError,TypeError,OverflowError):arrival=None;reason=reason or 'invalid_arrival'
        if not arrival:reason=reason or 'missing_arrival_anchor'
        elif arrival['day']>report['day']:reason=reason or 'future_arrival'
        departure=None;departure_error=None if departure_header else 'unverified_departure_header'
        if departure_header:
            try:departure=date_value(cells[8])
            except (ValueError,TypeError,OverflowError):departure_error='invalid_departure'
        if departure and (departure['day']>report['day'] or (arrival and departure['day']<arrival['day'])):
            departure_error='future_or_reversed_departure';departure=None
        record={'row':n,'scope':area,'vessel':name,'dwt':weight(cells[6]),'gt':weight(cells[5]),'arrival':arrival,'departure':departure,'departure_error':departure_error,'raw':evidence}
        if reason:rejected.append({**record,'reason':reason})
        else:rows.append(record)
    if not any(h['scope'] for h in headings):raise ValueError('No explicit NAP operator sections')
    if not valid_headers:raise ValueError('NAP column schema changed')
    return {'report_day':report['day'],'headings':headings,'schema_errors':schema_errors,'rows':rows,'rejected':rejected}

def collect(item,cached=False,refresh_after=None):
    folder=OUT/'parsed';folder.mkdir(exist_ok=True);path=folder/(str(item['id'])+'.json')
    if cached and path.exists():return json.loads(path.read_text(encoding='utf-8'))
    previous=json.loads(path.read_text(encoding='utf-8')) if path.exists() else None
    if refresh_after and item['day']<refresh_after and previous and previous.get('status')=='parsed' and not previous.get('source_error') and previous.get('parser_version')==VERSION:
        return {**previous,'historical_retained':True}
    result={**item,'parser_version':VERSION,'checked_at':dt.datetime.now(dt.timezone.utc).isoformat()}
    try:
        raw=fetch(item['url'],limit=500_000);result['html_sha256']=archive(raw,'.html')
        s=BeautifulSoup(raw,'html.parser');date=s.select_one('.dateNews')
        if not date or publication(date.get_text(' ',strip=True))[0]!=item['day']:raise ValueError('Detail/index publication mismatch')
        attachments=[a['href'] for a in s.select('a[href]') if re.search(r'\.xlsx(?:$|\?)',a['href'],re.I)]
        if len(attachments)!=1:raise ValueError('Expected exactly one XLSX attachment')
        url=urllib.parse.quote(urllib.parse.urljoin(BASE,attachments[0]),safe=':/%?=&')
        result['attachment_url']=url
        data=fetch(url,limit=500_000);result['sha256']=archive(data,'.xlsx')
        result.update(parse_workbook(data,item['day']));result['status']='parsed'
    except Exception as e:
        if previous and previous.get('status') in ('parsed','retained') and previous.get('parser_version')==VERSION and (OUT/'archive'/(previous.get('sha256','')+'.xlsx')).exists():
            result={**previous,'status':'retained','source_error':str(e),'failed_check_at':result['checked_at']}
        else:result.update(status='failed',error=str(e),rows=[],rejected=[])
    path.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    return result

def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--resume',action='store_true');p.add_argument('--cached',action='store_true');p.add_argument('--reparse',action='store_true');p.add_argument('--reparse-failed',action='store_true');p.add_argument('--refresh-recent',type=int);p.add_argument('--workers',type=int,default=3);p.add_argument('--from-date',default='2025-01-01');p.add_argument('--discover-only',action='store_true');a=p.parse_args()
    OUT.mkdir(parents=True,exist_ok=True)
    manifest=json.loads((OUT/'candidates.json').read_text(encoding='utf-8')) if a.cached or a.reparse or a.reparse_failed else discover(str(dt.date.fromisoformat(a.from_date)),a.resume)
    if a.discover_only:return
    results=[]
    if a.reparse or a.reparse_failed:
        for item in manifest['reports']:
            path=OUT/'parsed'/(str(item['id'])+'.json')
            report=json.loads(path.read_text(encoding='utf-8'))
            if report.get('sha256') and (not a.reparse_failed or report['status']!='parsed'):
                try:
                    report.update(parse_workbook((OUT/'archive'/(report['sha256']+'.xlsx')).read_bytes(),item['day']))
                    report['status']='parsed';report.pop('error',None)
                except Exception as e:report.update(status='failed',error=str(e),rows=[],rejected=[])
                path.write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
            results.append(report)
        result={'source':INDEX,'checked_at':manifest['checked_at'],'parser_version':VERSION,'reports':sorted(results,key=lambda r:(r['day'],r['published_at'],r['id']))}
        (OUT/'history.json').write_text(json.dumps(result,ensure_ascii=False),encoding='utf-8')
        print(json.dumps({'reports':len(results),'parsed':sum(r['status']=='parsed' for r in results),'rows':sum(len(r['rows']) for r in results),'failed':sum(r['status']=='failed' for r in results)}));return
    if a.refresh_recent is not None and not 1<=a.refresh_recent<=30:raise ValueError('Refresh window must be 1–30 days')
    refresh_after=(dt.date.fromisoformat(max(r['day'] for r in manifest['reports']))-dt.timedelta(days=a.refresh_recent)).isoformat() if a.refresh_recent else None
    with ThreadPoolExecutor(max_workers=max(1,min(a.workers,4))) as pool:
        jobs={pool.submit(collect,item,a.cached or a.resume,refresh_after):item for item in manifest['reports']}
        for job in as_completed(jobs):
            result=job.result();results.append(result)
            if len(results)%25==0 or result['status']=='failed':print(json.dumps({'completed':len(results),'status':result['status'],'id':result['id'],'error':result.get('error')}),flush=True)
    result={'source':INDEX,'checked_at':manifest['checked_at'],'parser_version':VERSION,'refresh_recent_days':a.refresh_recent,'reports':sorted(results,key=lambda r:(r['day'],r['published_at'],r['id']))}
    (OUT/'history.json').write_text(json.dumps(result,ensure_ascii=False),encoding='utf-8')
    print(json.dumps({'reports':len(results),'parsed':sum(r['status']=='parsed' for r in results),'rows':sum(len(r['rows']) for r in results),'failed':sum(r['status']=='failed' for r in results)}))

if __name__=='__main__':main()
